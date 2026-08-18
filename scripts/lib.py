"""Shared styling and helper utilities for the wedding planner workbook."""
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment, NamedStyle
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule, ColorScaleRule
from openpyxl.chart import BarChart, PieChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.drawing.fill import PatternFillProperties, ColorChoice
from openpyxl.chart.marker import DataPoint
from openpyxl.worksheet.pagebreak import Break
from openpyxl.utils.units import cm_to_EMU
import copy

# ---------------------------------------------------------------- PALETTE --
LAVENDER = "D8D3F0"
BLUSH = "F6D9DE"
SAGE = "DCE8D8"
CREAM = "F5EFD6"
POWDER = "D8E6F0"
PEACH = "F5DFD0"
CHARCOAL = "2E2E2E"
OFFWHITE = "FBF9F6"
GREY = "E5E1DA"
WHITE = "FFFFFF"

# darker accent variants used for banner text/fills where needed
DARK_LAVENDER = "B4ABDD"
DARK_BLUSH = "EFBCC6"
DARK_SAGE = "BFD6B8"
DARK_CREAM = "E9DDA0"
DARK_POWDER = "B7D3E6"
DARK_PEACH = "EFC5A9"

STATUS_RED = "F5C6C6"
STATUS_GREEN = "C9E4CA"
STATUS_YELLOW = "F5EFC6"

PALETTE_ROTATION = [LAVENDER, BLUSH, SAGE, CREAM, POWDER, PEACH]

# ---------------------------------------------------------------- FONTS ----
FONT_NAME = "Calibri"          # body font (universally available)
FONT_SERIF = "Georgia"         # editorial serif for titles (Win + Mac default)

def f_title(color=CHARCOAL, size=16):
    # Serif display face gives the workbook a soft, editorial "wedding" feel.
    return Font(name=FONT_SERIF, bold=False, size=size, color=color)

def f_eyebrow(color=CHARCOAL, size=8):
    return Font(name=FONT_NAME, bold=True, size=size, color=color)

def f_header(color=CHARCOAL, size=11):
    return Font(name=FONT_NAME, bold=True, size=size, color=color)

def f_body(size=10, color=CHARCOAL, italic=False, bold=False):
    return Font(name=FONT_NAME, size=size, color=color, italic=italic, bold=bold)

def f_kpi_number(size=18, color=CHARCOAL):
    return Font(name=FONT_NAME, bold=True, size=size, color=color)

def f_kpi_label(size=9, color=CHARCOAL):
    return Font(name=FONT_NAME, bold=True, size=size, color=color)

THIN_GREY = Side(style="thin", color=GREY)
BORDER_ALL = Border(left=THIN_GREY, right=THIN_GREY, top=THIN_GREY, bottom=THIN_GREY)
BORDER_BOTTOM = Border(bottom=THIN_GREY)

def fill(color):
    return PatternFill(fill_type="solid", start_color=color, end_color=color)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
LEFT_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)

# ------------------------------------------------------------ SHEET BASE ---
def new_sheet(wb, title, tab_color=LAVENDER):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab_color
    ws.sheet_view.zoomScale = 100
    return ws

def set_col_widths(ws, widths):
    """widths: dict of col_letter -> width"""
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

def title_banner(ws, text, color, row=1, col_start=1, col_end=10, height=28, size=16, font_color=CHARCOAL):
    ws.merge_cells(start_row=row, start_column=col_start, end_row=row, end_column=col_end)
    c = ws.cell(row=row, column=col_start, value=text)
    c.font = f_title(color=font_color, size=size)
    c.fill = fill(color)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = height
    for cc in range(col_start, col_end + 1):
        ws.cell(row=row, column=cc).fill = fill(color)
    return row + 1

def kpi_box(ws, row, col, label, value, color, width_cols=2, height=2, number_format=None, font_size=16):
    """Draws a small bordered KPI card starting at (row, col) spanning width_cols columns and `height` rows.
    Returns the cell holding the value (for later formula reference)."""
    r0, c0 = row, col
    r1, c1 = row + height - 1, col + width_cols - 1
    ws.merge_cells(start_row=r0, start_column=c0, end_row=r0, end_column=c1)
    lab = ws.cell(row=r0, column=c0, value=label)
    lab.font = f_kpi_label(color=CHARCOAL)
    lab.fill = fill(color)
    lab.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=r0 + 1, start_column=c0, end_row=r1, end_column=c1)
    val = ws.cell(row=r0 + 1, column=c0, value=value)
    val.font = f_kpi_number(size=font_size)
    val.fill = fill(WHITE)
    val.alignment = Alignment(horizontal="center", vertical="center")
    if number_format:
        val.number_format = number_format
    for rr in range(r0, r1 + 1):
        for cc in range(c0, c1 + 1):
            cell = ws.cell(row=rr, column=cc)
            cell.border = BORDER_ALL
    return val

def table_header(ws, row, col_start, headers, color, font_color=CHARCOAL):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=col_start + i, value=h)
        c.font = f_header(color=font_color)
        c.fill = fill(color)
        c.alignment = CENTER
        c.border = BORDER_ALL
    # Taller header row so two-line wrapped headers are never clipped.
    ws.row_dimensions[row].height = 30
    return row + 1

def style_data_row(ws, row, col_start, col_end, banding_color=None, base_color=WHITE):
    bg = banding_color if banding_color else base_color
    for cc in range(col_start, col_end + 1):
        c = ws.cell(row=row, column=cc)
        c.font = f_body()
        c.fill = fill(bg)
        c.border = BORDER_ALL
        c.alignment = LEFT

def band_table(ws, first_data_row, last_data_row, col_start, col_end, light_color):
    for r in range(first_data_row, last_data_row + 1):
        band = light_color if (r - first_data_row) % 2 == 1 else WHITE
        style_data_row(ws, r, col_start, col_end, banding_color=band)

def freeze_header(ws, row):
    # Freeze panes intentionally disabled: the frozen-pane divider draws a grey
    # line across the whole sheet in Excel's view, which reads as clutter on a
    # premium template. Kept as a no-op so callers don't need changing.
    return

def autosize(ws, col_min_widths):
    """col_min_widths: dict col_letter -> minimum width; sets widths (approx autosize)."""
    for col, w in col_min_widths.items():
        ws.column_dimensions[col].width = w

def add_dropdown(ws, cell_range, formula, allow_blank=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=allow_blank, showDropDown=False)
    ws.add_data_validation(dv)
    dv.add(cell_range)
    return dv

# Locale-independent checkbox glyphs stored as *text* (not booleans).
# Real Excel renders boolean TRUE/FALSE in the user's language (WAAR/ONWAAR in
# Dutch), and a number format does not override that. Storing the actual box
# glyph as text shows identically in every locale, in the cell and the dropdown.
CHK = "☑"    # ☑ checked
UNCHK = "☐"  # ☐ unchecked

def add_checkbox_col(ws, cell_range, default_unchecked=True):
    """Checkbox column backed by text glyphs (☑ / ☐), locale-independent.
    Empty cells are pre-filled with ☐ so the column always reads as checkboxes."""
    dv = DataValidation(type="list", formula1="=CheckBox", allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv)
    dv.add(cell_range)
    big = Font(name=FONT_NAME, size=13, color=CHARCOAL)
    for row in ws[cell_range]:
        for cell in row:
            if default_unchecked and cell.value in (None, "", False):
                cell.value = UNCHK
            elif cell.value is True:
                cell.value = CHK
            cell.font = big
            cell.alignment = CENTER
    return dv

def cf_status_colors(ws, cell_range, mapping):
    """mapping: dict of text value -> fill color hex, applies equal-text conditional format."""
    for text, color in mapping.items():
        rule = FormulaRule(formula=[f'EXACT({cell_range.split(":")[0]},"{text}")'], fill=fill(color))
        ws.conditional_formatting.add(cell_range, rule)

def cf_overdue(ws, date_range, status_cell_col_letter=None):
    """Highlight dates in the past red. date_range like 'C5:C50'"""
    first_cell = date_range.split(":")[0]
    rule = FormulaRule(formula=[f'AND({first_cell}<TODAY(),{first_cell}<>"")'], fill=fill(STATUS_RED))
    ws.conditional_formatting.add(date_range, rule)

def cf_bool_true(ws, cell_range, color=STATUS_GREEN):
    first_cell = cell_range.split(":")[0]
    rule = FormulaRule(formula=[f'{first_cell}=TRUE'], fill=fill(color))
    ws.conditional_formatting.add(cell_range, rule)

CHART_COLORS = [LAVENDER, BLUSH, SAGE, CREAM, POWDER, PEACH, DARK_LAVENDER, DARK_BLUSH, DARK_SAGE, DARK_CREAM]


def clean_chart_frame(chart, legend="r"):
    """Strip the default chart border/fill so charts sit cleanly on the sheet,
    remove gridlines, and place a tidy legend (or hide it with legend=None)."""
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties
    noline = GraphicalProperties()
    noline.noFill = True
    noline.line = LineProperties(noFill=True)
    chart.graphical_properties = noline
    # Plot data even when its source cells sit in hidden helper columns.
    chart.plotVisOnly = False
    chart.displayBlanksAs = "gap"
    try:
        pa = GraphicalProperties()
        pa.noFill = True
        pa.line = LineProperties(noFill=True)
        chart.plot_area.graphicalProperties = pa
    except Exception:
        pass
    if legend is None:
        chart.legend = None
    elif chart.legend is not None:
        chart.legend.position = legend
        chart.legend.overlay = False
    return chart

def style_chart_series_colors(chart, colors=None):
    from openpyxl.chart.marker import DataPoint
    from openpyxl.drawing.fill import PatternFillProperties
    colors = colors or CHART_COLORS
    for s in chart.series:
        s.graphicalProperties.line.noFill = True
        pts = []
        for i, color in enumerate(colors):
            dp = DataPoint(idx=i)
            dp.graphicalProperties.solidFill = color
            pts.append(dp)
        s.data_points = pts

def make_donut(ws, title, cat_ref, val_ref, anchor, colors=None, width=8, height=6,
               legend="b", hole=62):
    chart = DoughnutChart()
    chart.title = title if title else None
    chart.style = 10
    chart.add_data(val_ref, titles_from_data=False)
    chart.set_categories(cat_ref)
    chart.height = height
    chart.width = width
    style_chart_series_colors(chart, colors)
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showPercent = True
    chart.dataLabels.showCatName = False
    chart.dataLabels.showSerName = False
    chart.dataLabels.showVal = False
    chart.dataLabels.showLegendKey = False
    clean_chart_frame(chart, legend=legend)
    ws.add_chart(chart, anchor)
    return chart

def make_bar(ws, title, cat_ref, val_ref, anchor, colors=None, width=10, height=6, series_titles=None):
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = title
    chart.y_axis.majorGridlines = None
    chart.add_data(val_ref, titles_from_data=True)
    chart.set_categories(cat_ref)
    chart.height = height
    chart.width = width
    colors = colors or CHART_COLORS
    for i, s in enumerate(chart.series):
        s.graphicalProperties.solidFill = colors[i % len(colors)]
        s.graphicalProperties.line.noFill = True
    clean_chart_frame(chart, legend=("r" if series_titles else None))
    ws.add_chart(chart, anchor)
    return chart

def make_pie(ws, title, cat_ref, val_ref, anchor, colors=None, width=8, height=6, legend="b"):
    chart = PieChart()
    chart.title = title if title else None
    chart.style = 10
    chart.add_data(val_ref, titles_from_data=False)
    chart.set_categories(cat_ref)
    chart.height = height
    chart.width = width
    style_chart_series_colors(chart, colors)
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showPercent = True
    chart.dataLabels.showCatName = False
    chart.dataLabels.showSerName = False
    chart.dataLabels.showVal = False
    chart.dataLabels.showLegendKey = False
    clean_chart_frame(chart, legend=legend)
    ws.add_chart(chart, anchor)
    return chart

def hyperlink_cell(ws, row, col, text, target_sheet):
    c = ws.cell(row=row, column=col, value=text)
    c.hyperlink = f"#'{target_sheet}'!A1"
    c.font = Font(name=FONT_NAME, color="6B5CA5", underline="single", size=11)
    return c

import os as _os
_ASSET_DIR = _os.path.join(_os.path.dirname(__file__), "..", "assets")

def embed_image(ws, filename, anchor_cell, width_px, height_px):
    """Place a placeholder/photo PNG from assets/ at anchor_cell, sized in px."""
    from openpyxl.drawing.image import Image as XLImage
    path = _os.path.join(_ASSET_DIR, filename)
    if not _os.path.exists(path):
        return None
    img = XLImage(path)
    img.width = width_px
    img.height = height_px
    ws.add_image(img, anchor_cell)
    return img

CURRENCY_FMT = '$#,##0.00'
DATE_FMT = 'DD-MMM-YYYY'
PCT_FMT = '0%'
