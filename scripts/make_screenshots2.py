# -*- coding: utf-8 -*-
"""Renders additional tab screenshots (beyond the first 5) needed for the
Etsy marketing images: Vendor Selection, Venue Comparison, Food and Drinks,
Wedding Party, Stationery Checklist, Accommodation, Transportation,
Moodboard, Decor Inventory, Flower Arrangements, Attire and Makeup."""
import os, sys, subprocess
sys.path.insert(0, os.path.dirname(__file__))
from openpyxl import load_workbook
from openpyxl.worksheet.page import PageMargins

SRC = os.path.join(os.path.dirname(__file__), "..", "dist", "Wedding-Planner-All-In-One.xlsx")
TMP = os.path.join(os.path.dirname(__file__), "..", "dist", "_screenshot_src2.xlsx")
OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "qa_screens")
os.makedirs(OUT_DIR, exist_ok=True)

TARGETS = [
    ("Vendor Selection", "A1:N24"),
    ("Venue Comparison", "A1:F19"),
    ("Food and Drinks", "A1:M26"),
    ("Wedding Party", "A1:N24"),
    ("Stationery Checklist", "A1:R24"),
    ("Accommodation", "A1:E45"),
    ("Transportation", "A1:H24"),
    ("Moodboard", "A1:F19"),
    ("Decor Inventory", "A1:L24"),
    ("Flower Arrangements", "A1:L24"),
    ("Attire and Makeup", "A1:E24"),
]
DEPS = ["Get Started", "Guest List"]

wb = load_workbook(SRC)
keep = [t for t, _ in TARGETS] + DEPS
for name in list(wb.sheetnames):
    if name not in keep:
        del wb[name]

order = [t for t, _ in TARGETS] + DEPS
wb._sheets.sort(key=lambda s: order.index(s.title))

for name, area in TARGETS:
    ws = wb[name]
    ws.print_area = area
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.2, right=0.2, top=0.2, bottom=0.2, header=0, footer=0)

for name in DEPS:
    ws = wb[name]
    ws.print_area = "A1:A1"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

wb.active = 0
wb.save(TMP)
print("Saved trimmed workbook:", TMP)
print("Target order:", [t for t, _ in TARGETS])
