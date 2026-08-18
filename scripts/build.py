# -*- coding: utf-8 -*-
"""Builds Wedding-Planner-All-In-One.xlsx tab by tab."""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from lib import *
from data_lists import *

wb = Workbook()
wb.remove(wb.active)

OUT_PATH = os.path.join(os.path.dirname(__file__), "..", "dist", "Wedding-Planner-All-In-One.xlsx")

# =====================================================================
# TAB: LISTS (hidden) - dropdown source data + named ranges
# =====================================================================
def build_lists(wb):
    ws = new_sheet(wb, "Lists", tab_color=GREY)
    col = 1
    named = {}

    def write_list(header, values, col):
        ws.cell(row=1, column=col, value=header).font = f_header()
        for i, v in enumerate(values, start=2):
            ws.cell(row=i, column=col, value=v)
        letter = get_column_letter(col)
        rng = f"Lists!${letter}$2:${letter}${1+len(values)}"
        return rng

    lists_map = {
        "AssignedTo": ASSIGNED_TO,
        "Priority": PRIORITY,
        "YesNo": YES_NO,
        "YesNoAwaited": YES_NO_AWAITED,
        "GuestTag": GUEST_TAG,
        "GuestType": GUEST_TYPE,
        "MealPref": MEAL_PREF,
        "VendorCategory": VENDOR_CATEGORY,
        "AttireStatus": ATTIRE_STATUS,
        "WeddingPartyRole": WEDDING_PARTY_ROLE,
        "GiftStatus": GIFT_STATUS,
        "DecorLocation": DECOR_LOCATION,
        "DecorStatus": DECOR_STATUS,
        "BuyRent": BUY_RENT,
        "FloralStatus": FLORAL_STATUS,
        "AttireCategory": ATTIRE_CATEGORY,
        "AccommodationList": ACCOMMODATION_LIST,
        "StationeryCategory": STATIONERY_CATEGORY,
        "DesignStatus": DESIGN_STATUS,
        "PrintStatus": PRINT_STATUS,
        "PaymentStatus": PAYMENT_STATUS,
        "PaymentType": PAYMENT_TYPE,
        "TrueFalse": TRUE_FALSE,
        "CheckBox": CHECKBOX,
        "BachelorType": BACHELOR_TYPE,
        "CurrencyList": CURRENCY_LIST,
        "BookingStatus": BOOKING_STATUS,
        "MonthNames": MONTH_NAMES,
        "ChecklistSections": [s for s, _ in CHECKLIST_SECTIONS],
    }
    for header, values in lists_map.items():
        rng = write_list(header, values, col)
        named[header] = rng
        col += 1

    from openpyxl.workbook.defined_name import DefinedName
    for name, rng in named.items():
        wb.defined_names[name] = DefinedName(name, attr_text=rng)

    # Single-cell references holding the checkbox glyphs, so formulas can test
    # "is checked" without embedding the glyph (and its quotes) in a string.
    for glyph, gname in (("☑", "ChkVal"), ("☐", "UnchkVal")):
        ws.cell(row=1, column=col, value=glyph)
        letter = get_column_letter(col)
        wb.defined_names[gname] = DefinedName(gname, attr_text=f"Lists!${letter}$1")
        col += 1

    ws.column_dimensions["A"].width = 18
    return named

NAMED = build_lists(wb)

def dv_named(ws, cell_range, name, allow_blank=True):
    dv = DataValidation(type="list", formula1=f"={name}", allow_blank=allow_blank, showDropDown=False)
    ws.add_data_validation(dv)
    dv.add(cell_range)
    return dv

print("Lists tab built. Named ranges:", len(NAMED))

from tab_get_started import build_get_started_core, add_toc_links
GS_WS, GS_TOC_ROW = build_get_started_core(wb)
print("Get Started core built.")

from tab_checklist import build_checklist
CHECKLIST_META = build_checklist(wb, dv_named)
print("Checklist tab built:", CHECKLIST_META["total_items"], "items")

from tab_vendor_selection import build_vendor_selection
VENDOR_META = build_vendor_selection(wb, dv_named)
print("Vendor Selection tab built.")

from tab_budget import build_budget
BUDGET_META = build_budget(wb, dv_named, VENDOR_META)
print("Budget tab built.")

from tab_guest_list import build_guest_list
GUEST_META = build_guest_list(wb, dv_named)
print("Guest List tab built.")

from tab_seating import build_seating
RECEPTION_SEATING_META = build_seating(wb, dv_named, GUEST_META, "Reception Seating Plan", SAGE, n_tables=10, table_capacity=8)
print("Reception Seating tab built.")
REHEARSAL_SEATING_META = build_seating(wb, dv_named, GUEST_META, "Rehearsal Dinner Seating", PEACH, n_tables=4, table_capacity=8)
print("Rehearsal Dinner Seating tab built.")

from tab_group1 import build_itinerary, build_activities, build_music, build_packing
build_itinerary(wb)
build_activities(wb)
build_music(wb)
build_packing(wb, dv_named)
print("Group 1 tabs built (Itinerary, Activities, Music, Packing).")

from tab_group2 import build_venue_comparison, build_food_drinks, build_shot_list
build_venue_comparison(wb)
build_food_drinks(wb, dv_named)
build_shot_list(wb)
print("Group 2 tabs built (Venue Comparison, Food and Drinks, Shot List).")

from tab_group3 import build_registry, wire_registry, build_wedding_party, build_aisle_order, build_party_gifts
_reg_ws, _reg_first, _reg_last = build_registry(wb)
wire_registry(_reg_ws, _reg_first, _reg_last, dv_named)
WEDDING_PARTY_META = build_wedding_party(wb, dv_named)
build_aisle_order(wb)
PARTY_GIFTS_META = build_party_gifts(wb, dv_named)
print("Group 3 tabs built (Registry, Wedding Party, Aisle Order, Party Gifts).")

from tab_group4 import build_moodboard, build_decor, build_flowers, build_attire_makeup
build_moodboard(wb)
build_decor(wb, dv_named)
build_flowers(wb, dv_named)
build_attire_makeup(wb, dv_named)
print("Group 4 tabs built (Moodboard, Decor, Flowers, Attire and Makeup).")

from tab_group5 import build_accommodation, build_transportation, build_stationery, build_save_the_date
build_accommodation(wb, dv_named, GUEST_META)
build_transportation(wb, dv_named)
STATIONERY_META = build_stationery(wb, dv_named)
build_save_the_date(wb)
print("Group 5 tabs built (Accommodation, Transportation, Stationery, Save the Date).")

from tab_group6 import build_engagement_party, build_bridal_shower, build_bachelor
build_engagement_party(wb, dv_named)
build_bridal_shower(wb, dv_named)
build_bachelor(wb, dv_named)
print("Group 6 tabs built (Engagement Party, Bridal Shower, Bachelor(ette)).")

from tab_group7 import build_honeymoon, build_gifts_thank_you
build_honeymoon(wb, dv_named)
GIFTS_META = build_gifts_thank_you(wb, dv_named)
print("Group 7 tabs built (Honeymoon, Gifts and Thank You).")

from tab_smart_calendar import build_smart_calendar
build_smart_calendar(wb, dv_named, CHECKLIST_META, BUDGET_META)
print("Smart Calendar tab built.")

from tab_dashboard import build_dashboard
build_dashboard(wb, dv_named, CHECKLIST_META, BUDGET_META, GUEST_META, RECEPTION_SEATING_META, VENDOR_META)
print("Dashboard tab built.")

# ---- Finalize Get Started: table of contents with hyperlinks to every tab, hide Lists ----
TAB_ORDER = [
    "Dashboard", "Smart Calendar", "Wedding Checklist", "Wedding Itinerary", "Wedding Activities",
    "Music Planner", "Packing List", "Wedding Budget", "Vendor Selection", "Venue Comparison",
    "Food and Drinks", "Photo and Video Shot List", "Guest List", "Reception Seating Plan",
    "Rehearsal Dinner Seating", "Wedding Registry", "Wedding Party", "Aisle Order",
    "Wedding Party Gifts", "Moodboard", "Decor Inventory", "Flower Arrangements",
    "Attire and Makeup", "Accommodation", "Transportation", "Stationery Checklist",
    "Save the Date", "Engagement Party", "Bridal Shower", "Bachelor(ette) Planner",
    "Honeymoon Planner", "Gifts and Thank You",
]
add_toc_links(GS_WS, GS_TOC_ROW, TAB_ORDER)
print("Get Started table of contents linked.")

# Reorder sheets: Get Started first, Dashboard second, then the rest in TAB_ORDER, Lists hidden at the end
ORDERED_TITLES = ["Get Started"] + TAB_ORDER
wb._sheets.sort(key=lambda s: ORDERED_TITLES.index(s.title) if s.title in ORDERED_TITLES else len(ORDERED_TITLES))
LISTS_WS = wb["Lists"]
LISTS_WS.sheet_state = "hidden"
wb.active = 0

wb.save(OUT_PATH)
print("Saved to", OUT_PATH)
print("Total sheets:", len(wb.sheetnames))
