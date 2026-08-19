#!/usr/bin/env python3
"""
Build the "Calm Money Reset" tracker.

Takes the existing Budget Planner + Debt Tracker workbook
(scripts/source/BudgetTrackerEtsy.xlsx) and rebuilds it into the House of Calm
"Calm Money Reset" positioned version:

  * remaps the palette to the exact House of Calm pastel hexes
  * sets the Fraunces/Work Sans font system (Lora + Poppins + Work Sans in Sheets)
  * converts XLOOKUP to INDEX/MATCH so every formula computes in Excel,
    Google Sheets and LibreOffice alike
  * adds 5 new tabs: a Micro-Decision Dashboard, Doom Spending Log,
    Low-Spend Streak Tracker, Pause Before You Buy, and Monthly Reflection
  * keeps all existing locked-formula logic exactly as it was

Output: dist/Calm-Money-Reset-Tracker.xlsx
"""

import copy
import os
import warnings

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils import get_column_letter, column_index_from_string

warnings.filterwarnings("ignore")

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "source", "BudgetTrackerEtsy.xlsx")
OUT = os.path.join(HERE, "..", "dist", "Calm-Money-Reset-Tracker.xlsx")

# ---------------------------------------------------------------- palette ----
# House of Calm pastel system (exact spec hexes), openpyxl wants an alpha prefix
BLUSH = "FFF4C9CE"
SKY = "FFC9DDE8"
SAGE = "FFC8D8C4"
BUTTER = "FFF5E3B3"
LAV = "FFD9D0E8"
CREAM = "FFFBF6EF"
CHARCOAL = "FF3A3A3A"
WHITE = "FFFFFFFF"
MUTED = "FF8C8681"  # soft taupe for helper / secondary text

# soft tints for calm banding / card fills (kept gentle, never alarming)
BLUSH_SOFT = "FFFBE7E9"
SKY_SOFT = "FFE9F2F7"
SAGE_SOFT = "FFE7EFE4"
BUTTER_SOFT = "FFFBF2DA"
LAV_SOFT = "FFF0EBF7"

# old -> new fill remap (existing file was already near this palette)
FILL_MAP = {
    "FFFBF9F6": CREAM,
    "FFF5EBD2": BUTTER,
    "FFF7ECC9": BUTTER,
    "FFD6E8D6": SAGE,
    "FFCFE2ED": SKY,
    "FF3B3733": CHARCOAL,
    "FFF5D6DC": BLUSH,
    "FFE3DCF0": LAV,
    "FFE3DCEF": LAV,
}
FONT_COLOR_MAP = {
    "FF3B3733": CHARCOAL,
    "FFFBF9F6": CREAM,
    "FFB0AAA3": MUTED,
}
FONT_NAME_MAP = {
    "Aptos Narrow": "Work Sans",
    "Calibri": "Work Sans",
    "Georgia": "Lora",
}

TITLE = "Lora"      # Fraunces substitute for big titles (serif)
HEAD = "Poppins"    # section headers
BODY = "Work Sans"  # data / body

thin = Side(style="thin", color="FFEAE2D6")
NBORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def solid(hex8):
    return PatternFill(fill_type="solid", fgColor=hex8, bgColor=hex8)


def remap_existing_styles(wb):
    """Swap the old near-pastel palette for the exact House of Calm hexes and
    set the font system, without touching any formula or layout."""
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        for row in ws.iter_rows():
            for c in row:
                # fills
                f = c.fill
                if f is not None and f.patternType == "solid":
                    rgb = f.fgColor.rgb
                    if isinstance(rgb, str) and rgb in FILL_MAP:
                        c.fill = solid(FILL_MAP[rgb])
                # fonts
                if c.font is not None:
                    name = c.font.name
                    new_name = FONT_NAME_MAP.get(name, name)
                    col = c.font.color.rgb if (c.font.color and isinstance(c.font.color.rgb, str)) else None
                    new_col = FONT_COLOR_MAP.get(col, col)
                    if new_name != name or new_col != col:
                        nf = copy.copy(c.font)
                        nf.name = new_name
                        if new_col is not None:
                            nf.color = new_col
                        c.font = nf
        # recolour data-bar conditional formats to the exact palette sage
        for rng, rules in ws.conditional_formatting._cf_rules.items():
            for r in rules:
                if r.type == "dataBar" and r.dataBar is not None:
                    cur = r.dataBar.color.rgb if r.dataBar.color else None
                    r.dataBar.color = openpyxl.styles.Color(
                        rgb=FILL_MAP.get(cur, SAGE))


def sanitize_dashes(wb):
    """House style: no em-dashes in copy. Rewrites em/en dashes inside plain
    text cells (never formulas) to commas / hyphens that read naturally."""
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and not v.startswith("=") and (
                        "—" in v or "–" in v):
                    new = (v.replace(" — ", ", ")
                             .replace("—", ", ")
                             .replace("–", "-"))
                    while ", ," in new:
                        new = new.replace(", ,", ",")
                    while "  " in new:
                        new = new.replace("  ", " ")
                    c.value = new
                    n += 1
    return n


def convert_xlookup(wb):
    """XLOOKUP(v, CategoryList, CategoryTypeList) -> INDEX(CategoryTypeList,
    MATCH(v, CategoryList, 0)). Universal-compatible; also much faster to
    recalc. Handles both plain and array formulas."""
    import re

    pat = re.compile(
        r"_xlfn\.XLOOKUP\((?P<v>.+?),(?P<lookup>[A-Za-z0-9_]+),(?P<ret>[A-Za-z0-9_]+)\)"
    )

    def fix(text):
        return pat.sub(
            lambda m: f"INDEX({m.group('ret')},MATCH({m.group('v')},{m.group('lookup')},0))",
            text,
        )

    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, ArrayFormula):
                    if v.text and "XLOOKUP" in v.text:
                        c.value = ArrayFormula(v.ref, fix(v.text))
                        n += 1
                elif isinstance(v, str) and "XLOOKUP" in v:
                    c.value = fix(v)
                    n += 1
    return n


# --------------------------------------------------------- styling helpers ---
def banner(ws, cell_range, text, fill=BLUSH, font_color=CHARCOAL, size=22, height=42):
    """Full-width soft colour block title row."""
    first = cell_range.split(":")[0]
    ws.merge_cells(cell_range)
    c = ws[first]
    c.value = text
    c.fill = solid(fill)
    c.font = Font(name=TITLE, size=size, bold=True, color=font_color)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[int("".join(filter(str.isdigit, first)))].height = height
    # paint the rest of the merged band
    for rng_cell in ws[cell_range][0]:
        rng_cell.fill = solid(fill)


def subtitle(ws, cell_range, text):
    first = cell_range.split(":")[0]
    ws.merge_cells(cell_range)
    c = ws[first]
    c.value = text
    c.font = Font(name=BODY, size=10.5, italic=True, color=MUTED)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)


def section(ws, cell, text, fill=SAGE):
    c = ws[cell]
    c.value = text
    c.fill = solid(fill)
    c.font = Font(name=HEAD, size=12.5, bold=True, color=CHARCOAL)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)


def header_cell(ws, cell, text, fill=CHARCOAL, color=CREAM):
    c = ws[cell]
    c.value = text
    c.fill = solid(fill)
    c.font = Font(name=HEAD, size=10.5, bold=True, color=color)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = NBORDER


def label(ws, cell, text, bold=False, color=CHARCOAL, size=10.5, align="left"):
    c = ws[cell]
    c.value = text
    c.font = Font(name=BODY, size=size, bold=bold, color=color)
    c.alignment = Alignment(horizontal=align, vertical="center",
                            indent=1 if align == "left" else 0, wrap_text=True)


def input_cell(ws, cell, fill=CREAM, number_format=None, align="center"):
    c = ws[cell]
    c.fill = solid(fill)
    c.font = Font(name=BODY, size=10.5, color=CHARCOAL)
    c.alignment = Alignment(horizontal=align, vertical="center")
    c.border = NBORDER
    c.protection = Protection(locked=False)
    if number_format:
        c.number_format = number_format
    return c


def protect(ws):
    ws.protection.sheet = True
    ws.protection.selectLockedCells = False
    ws.protection.formatCells = False


CUR = '"$"#,##0.00'
CUR0 = '"$"#,##0'
PCT = "0%"


def build(wb):
    remap_existing_styles(wb)
    n = convert_xlookup(wb)
    print(f"  converted {n} XLOOKUP formulas -> INDEX/MATCH")
    d = sanitize_dashes(wb)
    print(f"  sanitized em-dashes in {d} legacy copy cells")

    add_settings_lists(wb)
    ws_dash = build_dashboard(wb)
    ws_doom = build_doom(wb)
    ws_streak = build_streak(wb)
    ws_pause = build_pause(wb)
    ws_reflect = build_reflection(wb)

    finalize(wb, ws_dash, ws_doom, ws_streak, ws_pause, ws_reflect)
    polish(wb)
    print("  polished legacy tabs for consistent widths / banners / spacing")


# canonical sheet names (emoji tab names; referenced only through these consts)
S_DASH = "Dashboard"
S_DOOM = "Doom Spending Log"
S_STREAK = "Low-Spend Streaks"
S_PAUSE = "Pause Before You Buy"
S_REFLECT = "Monthly Reflection"
FLEX = "Flexible Income Budget"
SND = "Savings & Debt Goal Tracker"

TAB_EMOJI = {
    "Welcome": "🌸 Welcome",
    S_DASH: "🌿 Dashboard",
    FLEX: FLEX,  # referenced by name -> leave unchanged
    "Transactions Log": "🪴 Transactions Log",
    "Bank Import": "🏦 Bank Import",
    S_DOOM: "💛 Doom Spending Log",
    S_STREAK: "🌱 Low-Spend Streaks",
    S_PAUSE: "✋ Pause Before You Buy",
    "Bill Calendar": "🗓️ Bill Calendar",
    SND: SND,  # scoped names live here -> leave unchanged
    "Subscription Tracker": "Subscription Tracker",  # formula-referenced
    "Habit Tracker": "🌼 Habit Tracker",
    S_REFLECT: "✨ Monthly Reflection",
    "Paycheck Budget": "Paycheck Budget",  # scoped names
    "Example": "Example",  # scoped names
    "Notes FAQ": "💬 Notes FAQ",
    "Settings": "Settings",  # global names + formula-referenced
}

TAB_COLORS = {
    "Welcome": "F4C9CE", S_DASH: "C8D8C4", FLEX: "C9DDE8",
    "Transactions Log": "F5E3B3", "Bank Import": "C9DDE8",
    S_DOOM: "F5E3B3", S_STREAK: "C8D8C4", S_PAUSE: "F4C9CE",
    "Bill Calendar": "D9D0E8", SND: "C8D8C4", "Subscription Tracker": "C9DDE8",
    "Habit Tracker": "F5E3B3", S_REFLECT: "D9D0E8", "Paycheck Budget": "C9DDE8",
    "Example": "FBF6EF", "Notes FAQ": "D9D0E8", "Settings": "FBF6EF",
}

ORDER = [
    "Welcome", S_DASH, FLEX, "Transactions Log", "Bank Import",
    S_DOOM, S_STREAK, S_PAUSE, "Bill Calendar", SND,
    "Subscription Tracker", "Habit Tracker", S_REFLECT,
    "Paycheck Budget", "Example", "Notes FAQ", "Settings",
]


def disp(name):
    """Final (display) sheet name, so cross-sheet formulas survive the rename."""
    return TAB_EMOJI.get(name, name)


def q(name):
    """Quote the final display sheet name for use in a formula."""
    return f"'{disp(name)}'"


def add_settings_lists(wb):
    ws = wb["Settings"]
    ws.protection.sheet = False  # unprotect to write, re-protect after
    lists = [
        ("V", "Doom Trigger", ["Boredom", "Stress", "Scrolling",
                               "Social pressure", "Celebration", "Other"]),
        ("X", "Doom Feeling", ["Fine", "Neutral", "Regret"]),
        ("Z", "Pause Decision", ["Bought", "Waited", "Skipped"]),
    ]
    for col, head, values in lists:
        hc = ws[f"{col}2"]
        hc.value = head
        hc.fill = solid(SKY)
        hc.font = Font(name=HEAD, size=10.5, bold=True, color=CHARCOAL)
        hc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for i, v in enumerate(values):
            c = ws[f"{col}{3 + i}"]
            c.value = v
            c.fill = solid(CREAM)
            c.font = Font(name=BODY, size=10.5, color=CHARCOAL)
            c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.column_dimensions[col].width = 16
    wb.defined_names.add(DefinedName("TriggerList", attr_text="Settings!$V$3:$V$8"))
    wb.defined_names.add(DefinedName("FeelingList", attr_text="Settings!$X$3:$X$5"))
    wb.defined_names.add(DefinedName("DecisionList", attr_text="Settings!$Z$3:$Z$5"))
    ws.protection.sheet = True


def _widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _add_dv(ws, formula1, sqref, allow_blank=True):
    dv = DataValidation(type="list", formula1=formula1, allow_blank=allow_blank,
                        sqref=sqref)
    dv.showErrorMessage = False
    ws.add_data_validation(dv)
    return dv


# ---------------------------------------------------------------- DASHBOARD ---
def build_dashboard(wb):
    ws = wb.create_sheet(S_DASH)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TAB_COLORS[S_DASH]
    _widths(ws, {"A": 3, "B": 13, "C": 13, "D": 13, "E": 13, "F": 13,
                 "G": 13, "H": 3, "I": 2, "J": 2})
    ws.column_dimensions["J"].hidden = True

    banner(ws, "A1:H1", "🌿 Calm Money Reset", fill=BLUSH, size=24, height=46)
    subtitle(ws, "A2:H2",
             "One calm glance. No dense charts, no red flags. Just what you "
             "need to decide right now.")
    ws.row_dimensions[2].height = 30

    def card(rng, lab, value, fill, numfmt, vfont=26):
        first = rng.split(":")[0]
        col = "".join(filter(str.isalpha, first))
        r = int("".join(filter(str.isdigit, first)))
        ws.merge_cells(rng)
        for cc in ws[rng][0]:
            cc.fill = solid(fill)
            cc.border = NBORDER
        # label on the card's first row, value below
        lc = ws[first]
        lc.value = lab
        lc.fill = solid(fill)
        lc.font = Font(name=HEAD, size=10.5, bold=True, color=CHARCOAL)
        lc.alignment = Alignment(horizontal="center", vertical="top")
        # value cell = one row down, merged across same cols
        vr = r + 1
        last_col = rng.split(":")[1]
        lcol = "".join(filter(str.isalpha, last_col))
        vrng = f"{col}{vr}:{lcol}{vr + 1}"
        ws.merge_cells(vrng)
        vc = ws[f"{col}{vr}"]
        vc.value = value
        vc.number_format = numfmt
        vc.font = Font(name=TITLE, size=vfont, bold=True, color=CHARCOAL)
        vc.alignment = Alignment(horizontal="center", vertical="center")
        for rr in range(vr, vr + 2):
            for cc in ws[f"{col}{rr}:{lcol}{rr}"][0]:
                cc.fill = solid(fill)
        # borders around whole card
        for rr in range(r, vr + 2):
            for cc in ws[f"{col}{rr}:{lcol}{rr}"][0]:
                cc.border = NBORDER

    # three big cards (rows 4-6 label+value)
    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 26
    ws.row_dimensions[6].height = 26
    card("B4:C4", "Safe to spend this week", None, SAGE_SOFT, CUR0)
    card("D4:E4", "Days until next bill", None, SKY_SOFT, '0" days"')
    card("F4:G4", "Low-spend streak", None, BLUSH_SOFT, '0" days"')

    ws["B5"] = (f"=IFERROR(MAX({q(FLEX)}!E10,0)/"
                f"MAX(ROUNDUP(({q(FLEX)}!B7-TODAY())/7,0),1),0)")
    ws["D5"] = ArrayFormula(
        "D5",
        '=IFERROR(IF(COUNTIFS(TransactionsLog[Type],"Bill",'
        'TransactionsLog[Date],">="&TODAY())=0,"None",'
        'MIN(IF((TransactionsLog[Type]="Bill")*'
        '(TransactionsLog[Date]>=TODAY()),TransactionsLog[Date],""))-TODAY()),'
        '"None")')
    ws["F5"] = f"=IFERROR({q(S_STREAK)}!$B$7,0)"

    # progress section
    section(ws, "A8", "Your progress", fill=SAGE)
    ws.merge_cells("A8:H8")
    ws.row_dimensions[8].height = 26

    # hidden pct helpers
    ws["J10"] = (f"=IFERROR(SUM({q(SND)}!E17:E24)/SUM({q(SND)}!B17:B24),0)")
    ws["J12"] = (f"=IFERROR(SUM({q(SND)}!E6:E13)/SUM({q(SND)}!B6:B13),0)")

    def bar(row, lab, pct_ref, emoji):
        label(ws, f"A{row}", f"{emoji}  {lab}", bold=True)
        ws.merge_cells(f"A{row}:C{row}")
        brng = f"D{row}:G{row}"
        ws.merge_cells(brng)
        bc = ws[f"D{row}"]
        bc.value = (f'=REPT("█",ROUND({pct_ref}*24,0))&'
                    f'REPT("░",24-ROUND({pct_ref}*24,0))')
        bc.font = Font(name=BODY, size=11, color="FF7BA87A")
        bc.alignment = Alignment(horizontal="left", vertical="center")
        pc = ws[f"H{row}"]
        pc.value = f"={pct_ref}"
        pc.number_format = PCT
        pc.font = Font(name=HEAD, size=12, bold=True, color=CHARCOAL)
        pc.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = 24

    bar(10, "Debt paid off", "$J$10", "💚")
    bar(12, "Savings goal", "$J$12", "🌷")

    # this-month strip
    section(ws, "A14", "This month, at a glance", fill=SKY)
    ws.merge_cells("A14:H14")
    ws.row_dimensions[14].height = 26
    strip = [
        ("B15", "Income", f"={q(FLEX)}!E5", SAGE_SOFT),
        ("D15", "Spent", f"={q(FLEX)}!O5", BLUSH_SOFT),
        ("F15", "Left to spend", f"={q(FLEX)}!E10", BUTTER_SOFT),
    ]
    for first, lab, val, fill in strip:
        col = "".join(filter(str.isalpha, first))
        endcol = chr(ord(col) + 1)
        ws.merge_cells(f"{col}15:{endcol}15")
        ws.merge_cells(f"{col}16:{endcol}16")
        lc = ws[f"{col}15"]
        lc.value = lab
        lc.font = Font(name=HEAD, size=10, bold=True, color=MUTED)
        lc.alignment = Alignment(horizontal="center", vertical="center")
        vc = ws[f"{col}16"]
        vc.value = val
        vc.number_format = CUR0
        vc.font = Font(name=TITLE, size=16, bold=True, color=CHARCOAL)
        vc.alignment = Alignment(horizontal="center", vertical="center")
        for rr in (15, 16):
            for cc in ws[f"{col}{rr}:{endcol}{rr}"][0]:
                cc.fill = solid(fill)
                cc.border = NBORDER
    ws.row_dimensions[15].height = 18
    ws.row_dimensions[16].height = 24

    # gentle affirmation
    ws.merge_cells("A18:H18")
    a = ws["A18"]
    a.value = ("This isn't about restriction. It's about awareness, and you're "
               "already here, paying attention. That counts. 🤍")
    a.fill = solid(LAV_SOFT)
    a.font = Font(name=BODY, size=10.5, italic=True, color=CHARCOAL)
    a.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[18].height = 34

    protect(ws)
    return ws


# ------------------------------------------------------------ DOOM SPEND LOG ---
def build_doom(wb):
    ws = wb.create_sheet(S_DOOM)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TAB_COLORS[S_DOOM]
    _widths(ws, {"A": 14, "B": 12, "C": 32, "D": 18, "E": 16, "F": 2,
                 "G": 3, "H": 18, "I": 10})

    banner(ws, "A1:E1", "💛 Doom Spending Log", fill=BUTTER, size=22)
    subtitle(ws, "A2:E2",
             "Just noticing, not judging. Log the little impulse buys so you "
             "can see your own patterns gently over time. No shame here.")
    ws.row_dimensions[2].height = 30

    # two soft stat cards
    for first, lab, val, fill in [
        ("A4", "Logged this month",
         "=COUNTIFS($A$9:$A$508,\">=\"&(EOMONTH(TODAY(),-1)+1),"
         "$A$9:$A$508,\"<=\"&EOMONTH(TODAY(),0))", SKY_SOFT),
        ("C4", "Spent this month",
         "=SUMIFS($B$9:$B$508,$A$9:$A$508,\">=\"&(EOMONTH(TODAY(),-1)+1),"
         "$A$9:$A$508,\"<=\"&EOMONTH(TODAY(),0))", BLUSH_SOFT),
    ]:
        col = "".join(filter(str.isalpha, first))
        endcol = chr(ord(col) + 1)
        ws.merge_cells(f"{col}4:{endcol}4")
        ws.merge_cells(f"{col}5:{endcol}5")
        lc = ws[f"{col}4"]
        lc.value = lab
        lc.font = Font(name=HEAD, size=10, bold=True, color=MUTED)
        lc.alignment = Alignment(horizontal="center", vertical="center")
        vc = ws[f"{col}5"]
        vc.value = val
        vc.number_format = CUR0 if "SUM" in val else "0"
        vc.font = Font(name=TITLE, size=18, bold=True, color=CHARCOAL)
        vc.alignment = Alignment(horizontal="center", vertical="center")
        for rr in (4, 5):
            for cc in ws[f"{col}{rr}:{endcol}{rr}"][0]:
                cc.fill = solid(fill)
                cc.border = NBORDER
    ws.row_dimensions[4].height = 18
    ws.row_dimensions[5].height = 26

    # by-trigger mini summary (right)
    ws["H4"] = "By trigger"
    ws["H4"].font = Font(name=HEAD, size=10.5, bold=True, color=CHARCOAL)
    triggers = ["Boredom", "Stress", "Scrolling", "Social pressure",
                "Celebration", "Other"]
    for i, t in enumerate(triggers):
        r = 5 + i
        ws[f"H{r}"] = t
        ws[f"H{r}"].font = Font(name=BODY, size=10, color=CHARCOAL)
        ws[f"H{r}"].alignment = Alignment(horizontal="left", vertical="center")
        ws[f"I{r}"] = f'=COUNTIF($D$9:$D$508,"{t}")'
        ws[f"I{r}"].font = Font(name=BODY, size=10, bold=True, color=CHARCOAL)
        ws[f"I{r}"].alignment = Alignment(horizontal="center", vertical="center")

    # table headers row 7
    heads = ["Date", "Amount", "What I bought", "Trigger", "How I feel now"]
    for i, h in enumerate(heads):
        header_cell(ws, f"{chr(65 + i)}7", h)
    ws.row_dimensions[7].height = 24
    ws.freeze_panes = "A8"

    # example row 8 (type over me)
    ex = ["=EOMONTH(TODAY(),-1)+3", 18.5, "Impulse phone case (example, type over me)",
          "Scrolling", "Neutral"]
    for i, v in enumerate(ex):
        c = ws[f"{chr(65 + i)}8"]
        c.value = v
        c.font = Font(name=BODY, size=10.5, italic=True, color=MUTED)
        c.alignment = Alignment(horizontal="center" if i in (0, 1, 3, 4) else "left",
                                vertical="center", indent=0 if i in (0, 1, 3, 4) else 1)
        c.border = NBORDER
        c.protection = Protection(locked=False)
    ws["A8"].number_format = "mmm d"
    ws["B8"].number_format = CUR

    # blank input rows 9-508
    last = 508
    for r in range(9, last + 1):
        for i in range(5):
            c = ws[f"{chr(65 + i)}{r}"]
            c.fill = solid(CREAM if r % 2 else SAGE_SOFT)
            c.font = Font(name=BODY, size=10.5, color=CHARCOAL)
            c.border = NBORDER
            c.protection = Protection(locked=False)
            c.alignment = Alignment(
                horizontal="center" if i in (0, 1, 3, 4) else "left",
                vertical="center", indent=0 if i in (0, 1, 3, 4) else 1)
        ws[f"A{r}"].number_format = "mmm d, yyyy"
        ws[f"B{r}"].number_format = CUR

    _add_dv(ws, "TriggerList", f"D8:D{last}")
    _add_dv(ws, "FeelingList", f"E8:E{last}")

    # soft conditional colour per trigger (gentle, never red)
    trig_fill = {"Boredom": BUTTER_SOFT, "Stress": SKY_SOFT, "Scrolling": LAV_SOFT,
                 "Social pressure": BLUSH_SOFT, "Celebration": SAGE_SOFT,
                 "Other": "FFF0EEEA"}
    for t, fill in trig_fill.items():
        ws.conditional_formatting.add(
            f"D8:D{last}",
            CellIsRule(operator="equal", formula=[f'"{t}"'], fill=solid(fill)))
    # feeling: regret gently in lavender (reflection cue, not alarm)
    ws.conditional_formatting.add(
        f"E8:E{last}",
        CellIsRule(operator="equal", formula=['"Regret"'], fill=solid(LAV_SOFT)))

    protect(ws)
    return ws


# --------------------------------------------------------- LOW-SPEND STREAK ---
def build_streak(wb):
    ws = wb.create_sheet(S_STREAK)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TAB_COLORS[S_STREAK]
    _widths(ws, {"A": 3, "B": 12, "C": 12, "D": 12, "E": 12, "F": 12, "G": 12,
                 "H": 12, "I": 3})
    for col in ("K", "L", "M", "N"):
        ws.column_dimensions[col].hidden = True

    banner(ws, "A1:H1", "Low-Spend Streak Tracker", fill=BLUSH, size=24, height=48)
    subtitle(ws, "A2:H2",
             "Every no-spend (or low-spend) day is a win. Mark a check and watch "
             "your streak grow. Miss a day? You just start a fresh one. That's it.")
    ws.row_dimensions[2].height = 40

    # month / year selectors (row 4)
    label(ws, "B4", "Month", bold=True, align="right")
    ws.merge_cells("C4:D4")
    input_cell(ws, "C4", fill=CREAM, align="left")
    ws["C4"] = "January"
    label(ws, "F4", "Year", bold=True, align="right")
    input_cell(ws, "G4", fill=CREAM, align="left")
    ws["G4"] = 2026
    _add_dv(ws, "MonthList", "C4")
    _add_dv(ws, "YearList", "G4")

    # helper cells (hidden)
    ws["L1"] = "=MATCH($C$4,MonthList,0)"
    ws["L2"] = "=DAY(EOMONTH(DATE($G$4,$L$1,1),0))"
    ws["L3"] = ("=IF(AND($L$1=MONTH(TODAY()),$G$4=YEAR(TODAY())),"
                "DAY(TODAY()),$L$2)")

    # two summary cards (rows 6-7), Dashboard style
    def streak_card(lrng, vrng, lab, fill):
        ws.merge_cells(lrng)
        lc = ws[lrng.split(":")[0]]
        lc.value = lab
        lc.font = Font(name=HEAD, size=11, bold=True, color=MUTED)
        lc.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(vrng)
        vc = ws[vrng.split(":")[0]]
        vc.font = Font(name=TITLE, size=26, bold=True, color=CHARCOAL)
        vc.alignment = Alignment(horizontal="center", vertical="center")
        vc.number_format = '0" days"'
        for rng in (lrng, vrng):
            for cc in ws[rng][0]:
                cc.fill = solid(fill)
                cc.border = NBORDER

    streak_card("B6:D6", "B7:D7", "CURRENT STREAK", BLUSH_SOFT)
    streak_card("F6:H6", "F7:H7", "LONGEST THIS MONTH", SAGE_SOFT)
    ws["B7"] = "=INDEX($M$6:$M$36,$L$3)"
    ws["F7"] = "=MAX($M$6:$M$36)"
    ws.row_dimensions[6].height = 20
    ws.row_dimensions[7].height = 40

    # calendar-style check grid: 5 label/check pairs, 7 wide (B..H)
    def check_coord(day):
        pair = (day - 1) // 7
        idx = (day - 1) % 7
        return chr(66 + idx), 11 + pair * 3   # (col, check row)

    section(ws, "A9", "Mark a check on every low-spend day", fill=SKY)
    ws.merge_cells("A9:H9")
    ws.row_dimensions[9].height = 28

    for pair in range(5):
        lrow = 10 + pair * 3
        crow = 11 + pair * 3
        for idx in range(7):
            day = pair * 7 + idx + 1
            col = chr(66 + idx)
            lc = ws[f"{col}{lrow}"]
            lc.value = f'=IF({day}<=$L$2,{day},"")'
            lc.font = Font(name=HEAD, size=10, bold=True, color=MUTED)
            lc.alignment = Alignment(horizontal="center", vertical="bottom")
            cc = ws[f"{col}{crow}"]
            cc.fill = solid(CREAM)
            cc.font = Font(name=BODY, size=16, bold=True, color="FF7BA87A")
            cc.alignment = Alignment(horizontal="center", vertical="center")
            cc.border = NBORDER
            cc.protection = Protection(locked=False)
        ws.row_dimensions[lrow].height = 18
        ws.row_dimensions[crow].height = 34

    check_cells = [f"{check_coord(day)[0]}{check_coord(day)[1]}"
                   for day in range(1, 32)]
    _add_dv(ws, '"✓, "', " ".join(check_cells))
    for cell in check_cells:
        ws.conditional_formatting.add(
            cell, CellIsRule(operator="equal", formula=['"✓"'], fill=solid(SAGE)))

    # streak helper table (hidden): rows 6..36 -> day = row-5
    for day in range(1, 32):
        r = 5 + day
        col, crow = check_coord(day)
        ws[f"K{r}"] = day
        ws[f"L{r}"] = f"={col}{crow}"
        prev = "0" if day == 1 else f"M{r - 1}"
        ws[f"M{r}"] = f'=IF(AND({day}<=$L$2,L{r}="✓"),{prev}+1,0)'

    # milestone micro-copy
    ws.merge_cells("A26:H26")
    m = ws["A26"]
    m.value = ('=IF($B$7>=30,"30 days. A whole month of showing up for '
               'yourself. Incredible.",IF($B$7>=14,"Two weeks strong. Your '
               'brain is learning a new pattern. Keep going.",IF($B$7>=7,'
               '"A full week. That is real momentum. Be proud.",IF($B$7>=1,'
               '"You have started a streak. One day at a time.",'
               '"A fresh page. Mark today whenever you are ready."))))')
    m.fill = solid(LAV_SOFT)
    m.font = Font(name=BODY, size=12, italic=True, color=CHARCOAL)
    m.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[26].height = 40

    protect(ws)
    return ws


# ------------------------------------------------------ PAUSE BEFORE YOU BUY ---
def build_pause(wb):
    ws = wb.create_sheet(S_PAUSE)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TAB_COLORS[S_PAUSE]
    _widths(ws, {"A": 24, "B": 12, "C": 16, "D": 16, "E": 28, "F": 14})

    banner(ws, "A1:F1", "✋ Pause Before You Buy", fill=BLUSH, size=22)
    subtitle(ws, "A2:F2",
             "The 24-hour rule, made simple. Before an impulse buy, log it here "
             "and give tomorrow-you a vote. Waiting is a win, not a loss.")
    ws.row_dimensions[2].height = 30

    last = 207
    # summary cards
    cards = [
        ("A4", "Times you paused", SAGE_SOFT,
         f'=COUNTIF($F$9:$F${last},"Waited")+COUNTIF($F$9:$F${last},"Skipped")',
         "0"),
        ("C4", "Money paused", BUTTER_SOFT,
         f'=SUMIF($F$9:$F${last},"Waited",$B$9:$B${last})+'
         f'SUMIF($F$9:$F${last},"Skipped",$B$9:$B${last})', CUR0),
        ("E4", "Bought anyway", BLUSH_SOFT,
         f'=COUNTIF($F$9:$F${last},"Bought")', "0"),
    ]
    for first, lab, fill, val, nf in cards:
        col = "".join(filter(str.isalpha, first))
        endcol = chr(ord(col) + 1)
        ws.merge_cells(f"{col}4:{endcol}4")
        ws.merge_cells(f"{col}5:{endcol}5")
        ws[f"{col}4"] = lab
        ws[f"{col}4"].font = Font(name=HEAD, size=10, bold=True, color=MUTED)
        ws[f"{col}4"].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"{col}5"] = val
        ws[f"{col}5"].number_format = nf
        ws[f"{col}5"].font = Font(name=TITLE, size=18, bold=True, color=CHARCOAL)
        ws[f"{col}5"].alignment = Alignment(horizontal="center", vertical="center")
        for rr in (4, 5):
            for cc in ws[f"{col}{rr}:{endcol}{rr}"][0]:
                cc.fill = solid(fill)
                cc.border = NBORDER
    ws.row_dimensions[4].height = 18
    ws.row_dimensions[5].height = 26

    heads = ["Item", "Cost", "Do I need this?", "Can it wait 24h?",
             "How will I feel tomorrow?", "Decision"]
    for i, h in enumerate(heads):
        header_cell(ws, f"{chr(65 + i)}7", h)
    ws.row_dimensions[7].height = 30
    ws.freeze_panes = "A8"

    # example row
    ex = ["Second pair of sneakers (example, type over me)", 74, "No", "Yes",
          "Relieved I waited", "Waited"]
    for i, v in enumerate(ex):
        c = ws[f"{chr(65 + i)}8"]
        c.value = v
        c.font = Font(name=BODY, size=10.5, italic=True, color=MUTED)
        c.border = NBORDER
        c.protection = Protection(locked=False)
        c.alignment = Alignment(
            horizontal="left" if i in (0, 4) else "center",
            vertical="center", indent=1 if i in (0, 4) else 0, wrap_text=True)
    ws["B8"].number_format = CUR

    for r in range(9, last + 1):
        for i in range(6):
            c = ws[f"{chr(65 + i)}{r}"]
            c.fill = solid(CREAM if r % 2 else SKY_SOFT)
            c.font = Font(name=BODY, size=10.5, color=CHARCOAL)
            c.border = NBORDER
            c.protection = Protection(locked=False)
            c.alignment = Alignment(
                horizontal="left" if i in (0, 4) else "center",
                vertical="center", indent=1 if i in (0, 4) else 0, wrap_text=True)
        ws[f"B{r}"].number_format = CUR

    _add_dv(ws, "YesNoList", f"C8:C{last}")
    _add_dv(ws, "YesNoList", f"D8:D{last}")
    _add_dv(ws, "DecisionList", f"F8:F{last}")

    ws.conditional_formatting.add(
        f"F8:F{last}",
        CellIsRule(operator="equal", formula=['"Waited"'], fill=solid(SAGE)))
    ws.conditional_formatting.add(
        f"F8:F{last}",
        CellIsRule(operator="equal", formula=['"Skipped"'], fill=solid(SKY)))
    ws.conditional_formatting.add(
        f"F8:F{last}",
        CellIsRule(operator="equal", formula=['"Bought"'], fill=solid(BUTTER)))

    protect(ws)
    return ws


# ---------------------------------------------------------- MONTHLY REFLECT ---
def build_reflection(wb):
    ws = wb.create_sheet(S_REFLECT)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TAB_COLORS[S_REFLECT]
    _widths(ws, {"A": 3, "B": 20, "C": 20, "D": 20, "E": 20, "F": 20, "G": 3})

    banner(ws, "A1:G1", "✨ Monthly Reflection", fill=LAV, size=22)
    subtitle(ws, "A2:G2",
             "Four gentle questions to close out the month. No scores, no "
             "grades, just a soft look back before you begin again.")
    ws.row_dimensions[2].height = 30

    label(ws, "B4", "Month", bold=True)
    input_cell(ws, "C4", fill=CREAM, align="left")
    _add_dv(ws, "MonthList", "C4")

    questions = [
        "1.  What felt good about my money this month?",
        "2.  What did I notice about my spending? (just noticing, no judgement)",
        "3.  What is one small, kind thing I want to try next month?",
        "4.  What am I proud of this month, money or otherwise?",
    ]
    row = 6
    for qtext in questions:
        section(ws, f"B{row}", qtext, fill=SAGE_SOFT)
        ws.merge_cells(f"B{row}:F{row}")
        ws.row_dimensions[row].height = 24
        # answer box
        arng = f"B{row + 1}:F{row + 2}"
        ws.merge_cells(arng)
        ac = ws[f"B{row + 1}"]
        ac.fill = solid(CREAM)
        ac.font = Font(name=BODY, size=11, color=CHARCOAL)
        ac.alignment = Alignment(horizontal="left", vertical="top",
                                 indent=1, wrap_text=True)
        ac.protection = Protection(locked=False)
        for rr in (row + 1, row + 2):
            for cc in ws[f"B{rr}:F{rr}"][0]:
                cc.fill = solid(CREAM)
                cc.border = NBORDER
        ws.row_dimensions[row + 1].height = 26
        ws.row_dimensions[row + 2].height = 26
        row += 4

    # Life Happens reset note
    ws.merge_cells(f"B{row}:F{row}")
    h = ws[f"B{row}"]
    h.value = "🤍  Life Happens"
    h.fill = solid(LAV)
    h.font = Font(name=HEAD, size=13, bold=True, color=CHARCOAL)
    h.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 28
    ws.merge_cells(f"B{row + 1}:F{row + 3}")
    note = ws[f"B{row + 1}"]
    note.value = ("If this month went sideways, you overspent, you stopped "
                  "logging, life got loud, that is completely okay. Nothing is "
                  "broken. You have not failed. Take a breath, start a new month "
                  "tab whenever you are ready, and pick up exactly where you are. "
                  "This tool waits for you, without judgement. Always.")
    note.fill = solid(LAV_SOFT)
    note.font = Font(name=BODY, size=11, italic=True, color=CHARCOAL)
    note.alignment = Alignment(horizontal="left", vertical="center",
                               indent=1, wrap_text=True)
    for rr in range(row + 1, row + 4):
        for cc in ws[f"B{rr}:F{rr}"][0]:
            cc.fill = solid(LAV_SOFT)
    ws.row_dimensions[row + 1].height = 28
    ws.row_dimensions[row + 2].height = 28
    ws.row_dimensions[row + 3].height = 28

    protect(ws)
    return ws


# -------------------------------------------------------------------- polish ---
import math as _math


def _spanw(widths, a, b):
    tot = 0
    for ci in range(column_index_from_string(a), column_index_from_string(b) + 1):
        tot += widths.get(get_column_letter(ci), 8.43)
    return tot


def _fit_h(text, wunits, size=10.5, minh=18):
    cpl = max(8, int(wunits * 0.92 * (11.0 / size)))
    lines = sum(max(1, _math.ceil(len(s) / cpl)) for s in str(text).split("\n"))
    return max(minh, lines * int(round(size * 1.65)) + 8)


def _rebanner(ws, span, title, height=46, size=23, fill=BLUSH):
    a = f"A1:{span}1"
    ws.merge_cells(a)
    c = ws["A1"]
    if title is not None:
        c.value = title
    c.fill = solid(fill)
    c.font = Font(name=TITLE, size=size, bold=True, color=CHARCOAL)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for cc in ws[a][0]:
        cc.fill = solid(fill)
    ws.row_dimensions[1].height = height


def _restyle_subtitle(ws, cellref, span, widths, size=10.5):
    """Style an instruction row as a wrapped muted subtitle and fit its height."""
    r = int("".join(filter(str.isdigit, cellref)))
    a0 = cellref
    a1 = f"{span}{r}"
    rng = f"{a0}:{a1}"
    # merge if not already
    already = any(str(m) == rng or (m.min_row == r and m.max_row == r)
                  for m in ws.merged_cells.ranges)
    if not already:
        try:
            ws.merge_cells(rng)
        except Exception:
            pass
    c = ws[a0]
    c.font = Font(name=BODY, size=size, italic=True, color=MUTED)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1,
                            wrap_text=True)
    col0 = "".join(filter(str.isalpha, a0))
    ws.row_dimensions[r].height = _fit_h(c.value or "",
                                         _spanw(widths, col0, span), size)


def _band(ws, cellref, span, widths, fill=SAGE, size=12.5):
    r = int("".join(filter(str.isdigit, cellref)))
    col0 = "".join(filter(str.isalpha, cellref))
    rng = f"{cellref}:{span}{r}"
    already = any(m.min_row == r and m.min_col == column_index_from_string(col0)
                  for m in ws.merged_cells.ranges)
    if not already and col0 != span:
        try:
            ws.merge_cells(rng)
        except Exception:
            pass
    c = ws[cellref]
    c.fill = solid(fill)
    c.font = Font(name=HEAD, size=size, bold=True, color=CHARCOAL)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for cc in ws[rng][0]:
        cc.fill = solid(fill)
    ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 0, 26)


BUD_W = {"A": 27, "B": 12, "C": 12, "D": 18, "E": 11, "F": 3, "G": 12, "H": 12,
         "I": 12, "J": 12, "K": 12, "L": 12, "M": 12}
# section titles that share their row with column headers (B..E): band A only
BUD_BANDS = [("A4", "B", SKY), ("D4", "E", SKY), ("A14", "A", SAGE),
             ("A23", "A", SAGE), ("A35", "A", SAGE), ("A47", "A", SAGE),
             ("A56", "A", SAGE), ("A65", "A", SAGE)]
BUD_WRAP = [(14, 30, "B", "E"), (23, 30, "B", "E"), (35, 30, "B", "E"),
            (47, 30, "B", "E"), (56, 30, "B", "E"), (70, 26, "A", "D")]

POLISH = {
    "🌸 Welcome": {"doc": True, "w": {"A": 116}, "banner": ("A", 52, 26),
                   "title": "Calm Money Reset"},
    "💬 Notes FAQ": {"doc": True, "w": {"A": 116}, "banner": ("A", 48, 24)},
    "Flexible Income Budget": {"w": BUD_W, "banner": ("F", 46, 22),
                               "sub": ("A2", "M"), "bands": BUD_BANDS,
                               "wrap_rows": BUD_WRAP, "note_rows": {69: "E"}},
    "Paycheck Budget": {"w": BUD_W, "banner": ("F", 46, 22),
                        "sub": ("A2", "M"), "bands": BUD_BANDS,
                        "wrap_rows": BUD_WRAP},
    "Example": {"w": BUD_W, "banner": ("F", 46, 22), "sub": ("A2", "M"),
                "bands": BUD_BANDS, "wrap_rows": BUD_WRAP},
    "🪴 Transactions Log": {"w": {"A": 14, "B": 24, "C": 14, "D": 14, "E": 40},
                            "banner": ("E", 46, 22), "sub": ("A2", "E")},
    "🏦 Bank Import": {"w": {"A": 15, "B": 34, "C": 14, "D": 22, "E": 14, "F": 16,
                            "G": 12, "H": 11, "I": 28, "J": 3, "K": 16, "L": 22},
                      "banner": ("I", 46, 22), "sub": ("A2", "I"),
                      "wrap_rows": [(5, 26, "A", "I")], "note_rows": {3: "I"}},
    "🗓️ Bill Calendar": {"w": {"A": 3, "B": 14, "C": 14, "D": 14, "E": 14,
                               "F": 14, "G": 14, "H": 14},
                         "banner": ("H", 46, 22), "sub": ("A2", "H"),
                         "hide": ["I", "J"]},
    "Savings & Debt Goal Tracker": {
        "w": {"A": 26, "B": 16, "C": 14, "D": 16, "E": 12, "F": 14, "G": 14},
        "banner": ("G", 46, 22), "sub": ("A2", "G"),
        "bands": [("A4", "G", SAGE), ("A15", "G", LAV), ("A27", "G", SKY)],
        "wrap_rows": [(5, 30, "A", "G"), (16, 30, "A", "G"), (30, 24, "A", "D")],
        "hide": ["I", "J"], "note_rows": {28: "G"}},
    "Subscription Tracker": {
        "w": {"A": 28, "B": 30, "C": 18, "D": 12, "E": 11, "F": 12, "G": 14,
              "H": 14, "I": 28, "K": 22, "L": 12},
        "banner": ("I", 46, 22), "sub": ("A2", "I"),
        "wrap_rows": [(6, 30, "A", "I")]},
    "🌼 Habit Tracker": {"w": {"A": 24, "B": 16, "AH": 14},
                        "banner": ("AG", 46, 22), "sub": ("A2", "AG")},
    "Settings": {"banner": ("A", 44, 20), "title": "Settings",
                 "sub": None, "note_rows": {1: "T"}},
}


def polish(wb):
    # rebrand the legacy product name in Welcome copy
    if "🌸 Welcome" in wb.sheetnames:
        w = wb["🌸 Welcome"]
        if w["A2"].value and "simple, automatic" in str(w["A2"].value):
            w["A2"] = ("An ADHD-friendly budget and debt tracker. Shame-free, "
                       "visual, and built to feel calm, not overwhelming.")

    for disp, cfg in POLISH.items():
        if disp not in wb.sheetnames:
            continue
        ws = wb[disp]
        ws.sheet_view.showGridLines = False
        widths = {}
        for k, v in cfg.get("w", {}).items():
            ws.column_dimensions[k].width = v
            widths[k] = v
        # fill remaining widths from existing for span math
        for col, cd in ws.column_dimensions.items():
            widths.setdefault(col, cd.width or 8.43)

        if "banner" in cfg:
            span, h, size = cfg["banner"]
            _rebanner(ws, span, cfg.get("title"), height=h, size=size)

        if cfg.get("sub"):
            cell, span = cfg["sub"]
            if ws[cell].value:
                _restyle_subtitle(ws, cell, span, widths)

        for col in cfg.get("hide", []):
            ws.column_dimensions[col].hidden = True

        for cell, span, fill in cfg.get("bands", []):
            if ws[cell].value is not None:
                _band(ws, cell, span, widths, fill=fill)

        for row, h, a, b in cfg.get("wrap_rows", []):
            for ci in range(column_index_from_string(a), column_index_from_string(b) + 1):
                c = ws.cell(row, ci)
                if c.value is not None:
                    al = c.alignment
                    c.alignment = Alignment(
                        horizontal=al.horizontal if al and al.horizontal else "center",
                        vertical="center", wrap_text=True,
                        indent=al.indent if al else 0)
            ws.row_dimensions[row].height = h

        for r, span in cfg.get("note_rows", {}).items():
            c = ws.cell(r, 1)
            if c.value:
                col0 = "A"
                rng = f"A{r}:{span}{r}"
                if not any(m.min_row == r and m.max_row == r
                           for m in ws.merged_cells.ranges):
                    try:
                        ws.merge_cells(rng)
                    except Exception:
                        pass
                c.alignment = Alignment(horizontal="left", vertical="center",
                                        indent=1, wrap_text=True)
                if not isinstance(c.value, str) or not c.value.startswith("="):
                    ws.row_dimensions[r].height = _fit_h(
                        c.value, _spanw(widths, col0, span))

        # document reflow: wrap + fit every text row in column A
        if cfg.get("doc"):
            wA = cfg["w"]["A"]
            for r in range(3, ws.max_row + 1):
                c = ws.cell(r, 1)
                if c.value is None or (isinstance(c.value, str) and
                                       c.value.startswith("=")):
                    continue
                al = c.alignment
                c.alignment = Alignment(horizontal="left", vertical="center",
                                        indent=1, wrap_text=True)
                # keep any existing section-band fill; just fit the height
                ws.row_dimensions[r].height = _fit_h(c.value, wA)


# ------------------------------------------------------------------ finalize ---
def finalize(wb, *sheets):
    # rename tabs (emoji) + tab colours
    for internal, display in TAB_EMOJI.items():
        if internal in wb.sheetnames and display != internal:
            wb[internal].title = display
    # tab colours by display name
    display_of = {k: v for k, v in TAB_EMOJI.items()}
    for internal, color in TAB_COLORS.items():
        disp = display_of.get(internal, internal)
        if disp in wb.sheetnames:
            wb[disp].sheet_properties.tabColor = color
    # reorder
    order_display = [display_of.get(n, n) for n in ORDER]
    present = [n for n in order_display if n in wb.sheetnames]
    present += [s.title for s in wb.worksheets if s.title not in present]
    wb._sheets.sort(key=lambda s: present.index(s.title))
    # open on the Dashboard
    wb.active = wb.sheetnames.index(display_of[S_DASH])
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False


if __name__ == "__main__":
    print("Loading source workbook ...")
    wb = openpyxl.load_workbook(SRC)
    build(wb)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print("Saved", os.path.abspath(OUT))
